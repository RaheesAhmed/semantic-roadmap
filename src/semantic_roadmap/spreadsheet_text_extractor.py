from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel


MAX_PREVIEW_ROWS_PER_SHEET = 50


class SpreadsheetSheetSummary(BaseModel):
    """Readable summary of one spreadsheet sheet."""

    sheet_name: str
    row_count: int
    column_count: int
    preview_rows: list[list[str]]


def extract_spreadsheet_summaries(
    spreadsheet_file_path: Path,
) -> list[SpreadsheetSheetSummary]:
    """Extract sheet structure and preview rows from an Excel workbook."""
    workbook = load_workbook(spreadsheet_file_path, read_only=True, data_only=True)
    sheet_summaries: list[SpreadsheetSheetSummary] = []

    for worksheet in workbook.worksheets:
        preview_rows: list[list[str]] = []
        for row_cells in worksheet.iter_rows(
            max_row=MAX_PREVIEW_ROWS_PER_SHEET,
            values_only=True,
        ):
            preview_rows.append(
                ["" if cell_value is None else str(cell_value) for cell_value in row_cells]
            )

        sheet_summaries.append(
            SpreadsheetSheetSummary(
                sheet_name=worksheet.title,
                row_count=worksheet.max_row,
                column_count=worksheet.max_column,
                preview_rows=preview_rows,
            )
        )

    workbook.close()
    return sheet_summaries

